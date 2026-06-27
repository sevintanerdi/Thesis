import pandas as pd
import numpy as np
import requests
from io import StringIO
from openpyxl import load_workbook

def get_ecb_data(series_key, start, end, freq='M'):
    url = f"https://data-api.ecb.europa.eu/service/data/{series_key}?startPeriod={start}&endPeriod={end}&format=csvdata"
    response = requests.get(url)
    df = pd.read_csv(StringIO(response.text))
    time_col = [c for c in df.columns if 'TIME' in c][0]
    val_col = [c for c in df.columns if 'OBS_VALUE' in c][0]
    df = df[[time_col, val_col]].copy()
    df.columns = ['date', 'value']
    df['value'] = pd.to_numeric(df['value'], errors='coerce')

    if freq == 'M':
        df['date'] = pd.to_datetime(df['date'])
        df = df.set_index('date').sort_index()
        df = df.resample('MS').mean().ffill()
        df = df.resample('QS').mean()
    else:
        df['date'] = df['date'].apply(
            lambda x: pd.to_datetime(f"{x[:4]}-{int(x[6])*3-2:02d}-01")
            if isinstance(x, str) and 'Q' in x else pd.NaT
        )
        df = df.set_index('date').sort_index()

    return df

def get_fred_data(series_id, col_name, start='2003-01-01'):
    url = f"https://fred.stlouisfed.org/graph/fredgraph.csv?id={series_id}"
    r = requests.get(url)
    df = pd.read_csv(StringIO(r.text))
    df.columns = ['date', 'value']
    df['date'] = pd.to_datetime(df['date'])
    df['value'] = pd.to_numeric(df['value'], errors='coerce')
    df = df.set_index('date').sort_index()
    df = df[start:]
    df = df.resample('QS').mean()
    df.columns = [col_name]
    return df

# ── ECB: Macro variables ──
policy_rate = get_ecb_data("FM/B.U2.EUR.4F.KR.DFR.LEV", "2003-01", "2025-12", "M")
policy_rate.columns = ['policy_rate']

inflation = get_ecb_data("ICP/M.U2.N.000000.4.ANR", "2003-01", "2025-12", "M")
inflation.columns = ['inflation']

gdp = get_ecb_data("MNA/Q.Y.I10.W2.S1.S1.B.B1GQ._Z._Z._Z.EUR.LR.G1", "2003-Q1", "2025-Q4", "Q")
gdp.columns = ['gdp_growth']

credit = get_ecb_data("BSI/M.U2.N.A.A20T.A.I.U2.2240.Z01.A", "2003-01", "2025-12", "M")
credit.columns = ['credit_growth']

# ── ECB: NIM - EA, DE, ES, IT, FR, NL ──
nim_ea = get_ecb_data("SUP/Q.B01.W0._Z.I2120._T.SII._Z._Z._Z.PCT.C", "2015-Q2", "2025-Q4", "Q")
nim_ea.columns = ['nim_ea']
nim_de = get_ecb_data("SUP/Q.DE.W0._Z.I2120._T.SII._Z._Z._Z.PCT.C", "2015-Q2", "2025-Q4", "Q")
nim_de.columns = ['nim_de']
nim_es = get_ecb_data("SUP/Q.ES.W0._Z.I2120._T.SII._Z._Z._Z.PCT.C", "2015-Q2", "2025-Q4", "Q")
nim_es.columns = ['nim_es']
nim_it = get_ecb_data("SUP/Q.IT.W0._Z.I2120._T.SII._Z._Z._Z.PCT.C", "2015-Q2", "2025-Q4", "Q")
nim_it.columns = ['nim_it']
nim_fr = get_ecb_data("SUP/Q.FR.W0._Z.I2120._T.SII._Z._Z._Z.PCT.C", "2015-Q2", "2025-Q4", "Q")
nim_fr.columns = ['nim_fr']
nim_nl = get_ecb_data("SUP/Q.NL.W0._Z.I2120._T.SII._Z._Z._Z.PCT.C", "2015-Q2", "2025-Q4", "Q")
nim_nl.columns = ['nim_nl']

# ── ECB: CET1 - EA, DE, ES, IT, FR, NL ──
cet1_ea = get_ecb_data("SUP/Q.B01.W0._Z.I4008._T.SII._Z._Z._Z.PCT.C", "2015-Q2", "2025-Q4", "Q")
cet1_ea.columns = ['cet1_ea']
cet1_de = get_ecb_data("SUP/Q.DE.W0._Z.I4008._T.SII._Z._Z._Z.PCT.C", "2015-Q2", "2025-Q4", "Q")
cet1_de.columns = ['cet1_de']
cet1_es = get_ecb_data("SUP/Q.ES.W0._Z.I4008._T.SII._Z._Z._Z.PCT.C", "2015-Q2", "2025-Q4", "Q")
cet1_es.columns = ['cet1_es']
cet1_it = get_ecb_data("SUP/Q.IT.W0._Z.I4008._T.SII._Z._Z._Z.PCT.C", "2015-Q2", "2025-Q4", "Q")
cet1_it.columns = ['cet1_it']
cet1_fr = get_ecb_data("SUP/Q.FR.W0._Z.I4008._T.SII._Z._Z._Z.PCT.C", "2015-Q2", "2025-Q4", "Q")
cet1_fr.columns = ['cet1_fr']
cet1_nl = get_ecb_data("SUP/Q.NL.W0._Z.I4008._T.SII._Z._Z._Z.PCT.C", "2015-Q2", "2025-Q4", "Q")
cet1_nl.columns = ['cet1_nl']

# ── FRED: 10Y Government Bond Yields ──
print("FRED'den bond yields çekiliyor...")
yield_de = get_fred_data("IRLTLT01DEM156N", "yield_de")
yield_es = get_fred_data("IRLTLT01ESM156N", "yield_es")
yield_it = get_fred_data("IRLTLT01ITM156N", "yield_it")
yield_fr = get_fred_data("IRLTLT01FRM156N", "yield_fr")
yield_nl = get_fred_data("IRLTLT01NLM156N", "yield_nl")

# ── Hepsini birleştir ──
all_dfs = [
    policy_rate, inflation, credit, gdp,
    nim_ea, nim_de, nim_es, nim_it, nim_fr, nim_nl,
    cet1_ea, cet1_de, cet1_es, cet1_it, cet1_fr, cet1_nl,
    yield_de, yield_es, yield_it, yield_fr, yield_nl
]
df = pd.concat(all_dfs, axis=1)
df = df.sort_index()
df = df['2015-04-01':]

# Manuel policy rate düzeltme
df.loc['2015-04-01', 'policy_rate'] = -0.20
df.loc['2015-07-01', 'policy_rate'] = -0.20
df['policy_rate'] = df['policy_rate'].ffill().bfill()

# ── SSR: Shadow Rate ──
wb = load_workbook('SSR_Estimates_20260602.xlsx', read_only=True, data_only=True)
ws = wb['D. Monthly average SSR series']
rows = list(ws.iter_rows(values_only=True))
ssr_data = [(row[1], row[3]) for row in rows[20:] if row[1] is not None and row[3] is not None]
df_ssr = pd.DataFrame(ssr_data, columns=['date', 'shadow_rate'])
df_ssr['date'] = pd.to_datetime(df_ssr['date'])
df_ssr = df_ssr.set_index('date').sort_index()
df_ssr_q = df_ssr['shadow_rate'].resample('QS').mean()
df = df.merge(df_ssr_q.rename('shadow_rate'), left_index=True, right_index=True, how='left')

# ── Kontrol ──
print("\n=== Her değişkenin gözlem sayısı ve son tarihi ===")
for col in df.columns:
    non_null = df[col].dropna()
    if len(non_null) > 0:
        print(f"{col}: {len(non_null)} obs, {non_null.index[0].date()} → {non_null.index[-1].date()}")
    else:
        print(f"{col}: NO DATA")

print("\n=== Eksik değerler ===")
print(df.isnull().sum())

df.to_csv('thesis_data.csv')
print("\nData kaydedildi!")